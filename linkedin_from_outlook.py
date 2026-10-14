"""
Excel "Email ID"  ->  Outlook contact card  ->  LinkedIn profile URL  ->  Excel "linkedin id"

Steps
  1. Ask the date (Enter = every row of this month).
  2. Read "redhat exam 2026.xlsx": rows with an Email ID but no linkedin id.
  3. Open Chrome with the "Arish - Chrome" profile at outlook.live.com/mail and
     click "New mail".
  4. For each email: put it in "To", hover over it so the contact card shows,
     click the LinkedIn logo, click "See profile on LinkedIn" and take the
     linkedin.com/in/... address from the tab that opens.
  5. Back up the Excel file and write the LinkedIn URLs into "linkedin id".
     The mail is never sent - the draft is discarded at the end.

Usage
  python linkedin_from_outlook.py                     # asks the date
  python linkedin_from_outlook.py --date 2026-10-07   # one day
  python linkedin_from_outlook.py --month 2026-10     # whole month
  python linkedin_from_outlook.py --limit 3           # only the first 3 emails (test)
  python linkedin_from_outlook.py --ask               # when a profile isn't found, wait so you can open it by hand
  python linkedin_from_outlook.py --dry-run           # don't save the Excel file
"""
import argparse
import datetime as dt
import json
import re
import shutil
import sys
import time
import urllib.parse
import urllib.request
from copy import copy
from pathlib import Path

import openpyxl
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys

import redhat_automation as pv
from redhat_automation import HERE, ask_user, click, click_text, log, visible

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.I)


# ---------------------------------------------------------------- date / excel

def ask_scope():
    """-> (date or None, any day of the month to use)."""
    today = dt.date.today()
    while True:
        text = input(f"\n>>> Which date? Type it as DD-MM-YYYY (e.g. {today:%d-%m-%Y}),\n"
                     f"    or just press Enter for all of {today:%B %Y}: ").strip(" \t﻿")
        if not text:
            print(f"    Using every row of {today:%B %Y}")
            return None, today
        try:
            date = pv.parse_date(text)
        except ValueError:
            print("    Not a valid date - try again, e.g. 07-10-2026")
            continue
        print(f"    Using {date:%A, %d-%b-%Y}")
        return date, date


def linkedin_col(ws, hr):
    for c in range(1, ws.max_column + 1):
        if re.sub(r"\s+", "", str(ws.cell(hr, c).value or "")).lower() == "linkedinid":
            return c
    raise RuntimeError(f"Column 'Linkedin Id' not found in sheet {ws.title}")


def rows_to_do(path, date, month_day):
    wb = openpyxl.load_workbook(path)
    ws = pv.month_sheet(wb, month_day)
    hr = pv.header_row(ws)
    cols = pv.sheet_columns(ws, hr)
    li = linkedin_col(ws, hr)
    todo = []
    # profiles already in the sheet (any day), to catch a stale contact card
    taken = {}
    for r in range(hr + 1, pv.last_data_row(ws, hr) + 1):
        url = linkedin_profile(str(ws.cell(r, li).value or ""))
        if url:
            taken[url.lower()] = pv.clean_cell_text(ws.cell(r, cols["name"]).value)
    for r in range(hr + 1, pv.last_data_row(ws, hr) + 1):
        if date and pv.cell_date(ws.cell(r, cols["date"]).value) != date:
            continue
        email = str(ws.cell(r, cols["email id"]).value or "").strip()
        if not EMAIL_RE.match(email) or not pv.is_blank(ws.cell(r, li).value):
            continue
        todo.append({"row": r, "name": pv.clean_cell_text(ws.cell(r, cols["name"]).value),
                     "email": email, "linkedin": ""})
    return ws.title, todo, taken


def save_results(path, sheet, results, backup_dir):
    found = [x for x in results if x["linkedin"]]
    if not found:
        log("No LinkedIn URLs found - Excel file not changed")
        return
    while True:
        try:
            with open(path, "r+b"):
                break
        except PermissionError:
            ask_user(f"The Excel file is open somewhere. Close it so the {len(found)} LinkedIn URL(s) can be saved.")
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    hr = pv.header_row(ws)
    cols = pv.sheet_columns(ws, hr)
    li = linkedin_col(ws, hr)
    written = 0
    for x in found:
        row = x["row"]
        # the row may have moved if someone edited the file meanwhile -> find it by email
        if str(ws.cell(row, cols["email id"]).value or "").strip().lower() != x["email"].lower():
            row = next((r for r in range(hr + 1, ws.max_row + 1)
                        if str(ws.cell(r, cols["email id"]).value or "").strip().lower() == x["email"].lower()), None)
        if not row or not pv.is_blank(ws.cell(row, li).value):
            continue
        cell = ws.cell(row, li)
        above = ws.cell(row - 1, li)
        if row - 1 > hr and above.has_style:
            cell._style = copy(above._style)
        cell.value = x["linkedin"]
        written += 1
        log(f"  row {row}: {x['name']} -> {x['linkedin']}")
    if not written:
        log("Nothing to write - Excel file not changed")
        return
    bdir = HERE / backup_dir
    bdir.mkdir(exist_ok=True)
    backup = bdir / f"{Path(path).stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.copy2(path, backup)
    log(f"Backup saved: {backup}")
    wb.save(path)
    log(f"Excel file saved - {written} LinkedIn URL(s) written to sheet {sheet}")


# ---------------------------------------------------------------- outlook

TO_BOX_CSS = ['div[contenteditable="true"][aria-label="To"]', 'input[aria-label="To"]',
              '[role="textbox"][aria-label^="To"]', 'div[contenteditable="true"][aria-label*="To"]']


def find_to_box(driver):
    for css in TO_BOX_CSS:
        for el in driver.find_elements(By.CSS_SELECTOR, css):
            if visible(el):
                return el
    return None


def open_compose(driver):
    end = time.time() + 40
    while time.time() < end:
        if find_to_box(driver):
            return find_to_box(driver)
        if "login" in driver.current_url or "signin" in driver.current_url:
            ask_user("Sign in to Outlook in the Arish Chrome window.")
            end = time.time() + 40
        elif click_text(driver, "New mail", "New message", "New email", timeout=3):
            log("Clicked New mail")
            time.sleep(3)
        else:
            time.sleep(1)
    ask_user("Could not open a new mail. Click 'New mail' in Outlook.")
    box = find_to_box(driver)
    if not box:
        raise RuntimeError("The 'To' box of the new mail was not found")
    return box


FIND_PILL_JS = r"""
const email = arguments[0].toLowerCase(), box = arguments[1];
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const has = e => [e.getAttribute('aria-label'), e.getAttribute('title'), e.innerText]
                 .some(v => (v || '').toLowerCase().includes(email));
// the recipient "pill" sits next to the To box - search the compose header around it
let area = box;
for (let i = 0; i < 6 && area.parentElement; i++) area = area.parentElement;
const hits = [...area.querySelectorAll('*')].filter(e => vis(e) && e !== box && has(e) && !e.isContentEditable);
// innermost element that matches
return hits.filter(e => !hits.some(o => o !== e && e.contains(o)))[0] || null;
"""

FIND_LINKEDIN_JS = r"""
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const label = e => [e.getAttribute('aria-label'), e.getAttribute('title'), e.getAttribute('alt'),
                    e.getAttribute('data-icon-name'), e.innerText].join(' ');
const els = [...document.querySelectorAll('button,a,[role=button],[role=tab],[role=link],img,svg,i')]
            .filter(e => vis(e) && /linked\s*in/i.test(label(e)) && !/see\s+profile|view\s+(full\s+)?profile/i.test(label(e)));
if (!els.length) return null;
const e = els[0];
return e.closest('button,a,[role=button],[role=tab],[role=link]') || e;
"""

FIND_SEE_PROFILE_JS = r"""
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const els = [...document.querySelectorAll('a,button,[role=button],[role=link]')].filter(e => vis(e) &&
    /(see|view)\s+(full\s+)?profile\s+on\s+linked\s*in|profile\s+on\s+linked\s*in/i.test(
        (e.innerText || '') + ' ' + (e.getAttribute('aria-label') || '') + ' ' + (e.getAttribute('title') || '')));
return els[0] || null;
"""


def linkedin_profile(url):
    """Any URL (also LinkedIn redirect / sign-in links) -> https://www.linkedin.com/in/<id>/"""
    for _ in range(3):
        m = re.search(r"linkedin\.com/in/([^/?#&\s\"']+)", url or "", re.I)
        if m:
            return f"https://www.linkedin.com/in/{m.group(1)}/"
        url = urllib.parse.unquote(url or "")
    return ""


def wait_for(driver, js, *args, timeout=8):
    end = time.time() + timeout
    while True:
        el = driver.execute_script(js, *args)
        if el or time.time() > end:
            return el
        time.sleep(0.5)


def hover(driver, el):
    """Put the mouse on the element, then move it slightly (the card shows on a light move)."""
    ac = ActionChains(driver)
    ac.move_to_element(el).pause(0.4).move_by_offset(3, 0).pause(0.4).move_by_offset(-3, 1).perform()


def new_linkedin_tab(port, before, timeout=20):
    """Wait for a tab that wasn't open before to show a LinkedIn address; return its URL and close it."""
    end = time.time() + timeout
    seen = None
    while time.time() < end:
        for t in pv.chrome_pages(port):
            if t["id"] in before:
                continue
            seen = t
            if linkedin_profile(t["url"]):
                close_tab(port, t["id"])
                return t["url"]
        time.sleep(0.5)
    if seen:
        log(f"    new tab did not reach a profile: {seen['url'][:100]}")
        close_tab(port, seen["id"])
    return ""


def close_tab(port, tab_id):
    try:
        urllib.request.urlopen(f"http://127.0.0.1:{port}/json/close/{tab_id}", timeout=10).close()
    except Exception:
        pass


def clear_to(driver):
    box = find_to_box(driver)
    if not box:
        return
    click(driver, box)
    for _ in range(4):  # Backspace in an empty To box removes the recipient before it
        box.send_keys(Keys.BACKSPACE)
        time.sleep(0.2)


def close_card(driver):
    box = find_to_box(driver)
    if box:
        ActionChains(driver).move_to_element(box).perform()  # mouse away from the pill
    if wait_for(driver, FIND_LINKEDIN_JS, timeout=0) or wait_for(driver, FIND_SEE_PROFILE_JS, timeout=0):
        ActionChains(driver).send_keys(Keys.ESCAPE).perform()
        time.sleep(0.5)


def lookup(driver, port, outlook_tab, email):
    """Email -> LinkedIn profile URL ('' if Outlook shows none)."""
    box = find_to_box(driver) or open_compose(driver)
    clear_to(driver)
    box = find_to_box(driver)
    click(driver, box)
    box.send_keys(email)
    time.sleep(1.5)
    box.send_keys(";")  # turns the typed address into a recipient
    pill = wait_for(driver, FIND_PILL_JS, email, box, timeout=8)
    if not pill:
        log("    recipient not found in the To box")
        return ""

    hover(driver, pill)
    li = wait_for(driver, FIND_LINKEDIN_JS, timeout=6)
    if not li:  # hover again, then try a click on the recipient
        hover(driver, pill)
        li = wait_for(driver, FIND_LINKEDIN_JS, timeout=4)
    if not li:
        click(driver, pill)
        li = wait_for(driver, FIND_LINKEDIN_JS, timeout=6)
    if not li:
        log("    contact card has no LinkedIn logo")
        return ""
    click(driver, li)

    see = wait_for(driver, FIND_SEE_PROFILE_JS, timeout=10)
    if not see:
        log("    no 'See profile on LinkedIn' (no LinkedIn match for this email?)")
        return ""
    direct = linkedin_profile(see.get_attribute("href") or "")
    before = {t["id"] for t in pv.chrome_pages(port)}
    click(driver, see)
    url = new_linkedin_tab(port, before) or direct
    driver.switch_to.window(outlook_tab)
    return linkedin_profile(url)


def manual_lookup(driver, port, outlook_tab, person):
    before = {t["id"] for t in pv.chrome_pages(port)}
    ask_user(f"Open the LinkedIn profile of {person['name']} ({person['email']}) by hand "
             "(hover -> LinkedIn logo -> See profile on LinkedIn).\n"
             "    Leave the LinkedIn tab open. Just press Enter to skip.")
    url = ""
    for t in pv.chrome_pages(port):
        if t["id"] not in before and linkedin_profile(t["url"]):
            url = t["url"]
            close_tab(port, t["id"])
    driver.switch_to.window(outlook_tab)
    return linkedin_profile(url)


def discard_draft(driver):
    clear_to(driver)
    # the Discard button is an icon - find it by its label
    btn = next((e for e in driver.find_elements(By.CSS_SELECTOR, '[aria-label="Discard"]') if visible(e)), None)
    if btn:
        click(driver, btn)
    elif not click_text(driver, "Discard", timeout=3):
        log("Discard button not found - delete the empty draft from Outlook Drafts")
        return
    time.sleep(1)
    click_text(driver, "OK", "Discard", "Yes", timeout=3)  # confirmation, if Outlook asks
    log("Draft discarded")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", help="YYYY-MM-DD: only that day")
    ap.add_argument("--month", help="YYYY-MM: every row of that month")
    ap.add_argument("--limit", type=int, help="only the first N emails")
    ap.add_argument("--ask", action="store_true", help="wait for you when a profile isn't found")
    ap.add_argument("--dry-run", action="store_true", help="don't save the Excel file")
    args = ap.parse_args()

    cfg = pv.load_config()
    if args.date:
        date = month_day = dt.date.fromisoformat(args.date)
    elif args.month:
        date, month_day = None, dt.date.fromisoformat(args.month + "-01")
    else:
        date, month_day = ask_scope()
    run(cfg, date, month_day, limit=args.limit, ask=args.ask, dry_run=args.dry_run)
    log("Done")


def run(cfg, date, month_day, limit=None, ask=False, dry_run=False):
    """Find LinkedIn URLs for the rows of `date` (None = the whole month of month_day)."""
    excel = cfg.get("excel", "path")
    pv.check_excel_closed(excel)

    sheet, todo, taken = rows_to_do(excel, date, month_day)
    if limit:
        todo = todo[:limit]
    log(f"Sheet {sheet}: {len(todo)} email(s) without a LinkedIn Id"
        + (f" on {date:%d-%b-%Y}" if date else ""))
    if not todo:
        return
    for x in todo:
        log(f"  row {x['row']}: {x['name']} | {x['email']}")

    port = cfg.getint("chrome", "debug_port")

    def start_outlook():
        d = pv.open_chrome(cfg, "linkedin", cfg.get("linkedin", "outlook_url"), "live.com")
        open_compose(d)
        return d, d.current_window_handle

    def back_to_outlook():
        """Return to the Outlook tab; if it was closed, open Outlook again."""
        nonlocal driver, outlook_tab
        try:
            driver.switch_to.window(outlook_tab)
            driver.execute_script("return 1")
        except Exception:
            log("    Outlook tab was closed - opening it again")
            driver, outlook_tab = start_outlook()

    driver, outlook_tab = start_outlook()

    logs = HERE / "logs"
    logs.mkdir(exist_ok=True)
    progress = logs / f"linkedin_{dt.datetime.now():%Y%m%d_%H%M%S}.json"
    try:
        for i, x in enumerate(todo, 1):
            log(f"[{i}/{len(todo)}] {x['name']} | {x['email']}")
            try:
                x["linkedin"] = lookup(driver, port, outlook_tab, x["email"])
            except Exception as e:
                log(f"    error: {type(e).__name__}: {str(e).splitlines()[0][:150]}")
                back_to_outlook()
            if not x["linkedin"] and ask:
                x["linkedin"] = manual_lookup(driver, port, outlook_tab, x)
            owner = taken.get(x["linkedin"].lower()) if x["linkedin"] else None
            if owner:
                # Outlook kept showing the card of an earlier person
                log(f"    ! got {x['linkedin']} which is {owner}'s profile - stale contact card, skipped")
                x["linkedin"] = ""
            elif x["linkedin"]:
                taken[x["linkedin"].lower()] = x["name"]
            log(f"    LinkedIn: {x['linkedin'] or '(not found)'}")
            try:
                close_card(driver)
            except Exception:
                back_to_outlook()
            progress.write_text(json.dumps(todo, indent=2), encoding="utf-8")
            time.sleep(1)
    except KeyboardInterrupt:
        log("Stopped - saving what was found so far")
    finally:
        try:
            discard_draft(driver)
        except Exception:
            log("Could not discard the draft - delete it from Outlook Drafts")
        found = sum(1 for x in todo if x["linkedin"])
        log(f"LinkedIn found for {found} of {len(todo)} email(s)")
        for x in todo:
            if not x["linkedin"]:
                log(f"  ! not found: {x['name']} | {x['email']}")
        if dry_run:
            log("Dry run - Excel file NOT saved")
        else:
            save_results(excel, sheet, todo, cfg.get("excel", "backup_dir"))


if __name__ == "__main__":
    main()
