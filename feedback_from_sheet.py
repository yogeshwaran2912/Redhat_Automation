"""
ADV PRO FEEDBACK (Responses) Google Sheet  ->  "redhat exam 2026.xlsx"

Steps
  1. Open Chrome (Ranjit profile), Google Drive, the "ADV PRO FEEDBACK (Responses)"
     sheet, and take that day's responses.
  2. Find each response's candidate in the month sheet (by phone or email, else by
     name on that day; then by phone or email anywhere in the month).
  3. Fill: Company Name <- "Your Company name", Designation <- "Your Designation",
     Advantage Pro Feedback form <- "y", Remark <- "Let others know about your
     experience...". Cells someone already filled are left alone.
  4. Back up the Excel file and save.

Responses that match nobody in the Red Hat sheet are listed at the end
(most are Pearson VUE candidates).

Usage
  python feedback_from_sheet.py                    # asks for the date (Enter = today)
  python feedback_from_sheet.py --date 2026-10-07  # a specific day
  python feedback_from_sheet.py --dry-run          # don't save the Excel file
"""
import argparse
import csv
import datetime as dt
import re
import shutil
import time
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

from redhat_automation import (HERE, CONFIG_PATH, ask_date, ask_user, cell_date, chrome_pages, activate_tab,
                               check_excel_closed, clean_cell_text, header_row, is_blank, last_data_row, load_config,
                               log, month_sheet, open_chrome, phone10)


# ---------------------------------------------------------------- google sheet

def save_config_value(section, key, value):
    """Update one key in config.ini without losing the comments."""
    lines = CONFIG_PATH.read_text(encoding="utf-8").splitlines()
    current = None
    for i, line in enumerate(lines):
        s = line.strip()
        if s.startswith("[") and s.endswith("]"):
            current = s[1:-1]
        elif current == section and re.match(rf"{key}\s*=", s):
            lines[i] = f"{key} = {value}"
            break
    CONFIG_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def find_feedback_sheet(driver, cfg):
    url = cfg.get("feedback", "sheet_url", fallback="").strip()
    if url:
        return url
    driver.get("https://drive.google.com/drive/search?q=ADV%20PRO%20FEEDBACK%20(Responses)")
    time.sleep(3)
    if "accounts.google.com" in driver.current_url:
        ask_user("Sign in to advpromktg@gmail.com in the Chrome window (only needed once).")
        driver.get("https://drive.google.com/drive/search?q=ADV%20PRO%20FEEDBACK%20(Responses)")
    ask_user("Double-click 'ADV PRO FEEDBACK (Responses)' in Google Drive so the sheet opens.")
    port = cfg.getint("chrome", "debug_port")
    for tab in chrome_pages(port):
        if "/spreadsheets/d/" in tab["url"]:
            activate_tab(port, tab["id"])
            driver.switch_to.window(tab["id"])
            save_config_value("feedback", "sheet_url", tab["url"])
            log("Saved the feedback sheet link in config.ini")
            return tab["url"]
    raise RuntimeError("The feedback sheet was not opened")


def download_feedback_csv(driver, sheet_url):
    m = re.search(r"/spreadsheets/d/([\w-]+)", sheet_url)
    gid = re.search(r"[#&?]gid=(\d+)", sheet_url)
    export = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    if gid:
        export += f"&gid={gid.group(1)}"
    dl = HERE / "downloads"
    dl.mkdir(exist_ok=True)
    for f in dl.glob("*.csv"):
        f.unlink()
    try:
        driver.execute_cdp_cmd("Browser.setDownloadBehavior", {"behavior": "allow", "downloadPath": str(dl)})
    except Exception:
        driver.execute_cdp_cmd("Page.setDownloadBehavior", {"behavior": "allow", "downloadPath": str(dl)})
    driver.get(export)
    for _ in range(60):
        files = list(dl.glob("*.csv"))
        if files and not list(dl.glob("*.crdownload")):
            return files[0]
        if "denied" in (driver.title or "").lower():  # Google sometimes refuses the export link
            break
        time.sleep(1)
    return fetch_feedback_csv(driver, sheet_url, m.group(1), gid.group(1) if gid else "", dl)


FETCH_JS = """const cb = arguments[arguments.length - 1];
fetch(arguments[0], {credentials: 'include'})
  .then(async r => cb({status: r.status, type: r.headers.get('content-type') || '', text: await r.text()}))
  .catch(e => cb({error: String(e)}));"""


def fetch_feedback_csv(driver, sheet_url, sheet_id, gid, dl):
    """Fallback: read the CSV from inside the open sheet (Google's 'gviz' link)."""
    log("Export link refused - reading the sheet another way "
        "(phone numbers typed with spaces come out blank; email and name still match)")
    driver.get(sheet_url)
    time.sleep(3)
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv" + (f"&gid={gid}" if gid else "")
    driver.set_script_timeout(60)
    r = driver.execute_async_script(FETCH_JS, url)
    if r.get("error") or r.get("status") != 200 or "csv" not in r.get("type", ""):
        raise RuntimeError("Feedback sheet could not be read (is Chrome signed in to advpromktg@gmail.com?) "
                           f"- {r.get('error') or r.get('status')}")
    path = dl / "feedback.csv"
    path.write_text(r["text"], encoding="utf-8")
    return path


def timestamp_dates(text, formats=None):
    text = text.strip()
    formats = formats or ["%m/%d/%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S",
                          "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d"]
    out = set()
    for f in formats:
        try:
            out.add(dt.datetime.strptime(text, f).date())
        except ValueError:
            pass
    return out


def detect_ts_format(values):
    """Google exports dates in the sheet's locale: decide month/day vs day/month
    from values like 10/27/2026 or 27/10/2026, so 7/10/2026 isn't read both ways."""
    mdy = dmy = 0
    for v in values:
        m = re.match(r"\s*(\d{1,2})/(\d{1,2})/\d{4}", v)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            mdy += b > 12
            dmy += a > 12
    if mdy == dmy:
        return ""
    return "%m/%d/%Y" if mdy > dmy else "%d/%m/%Y"


def read_feedback(csv_path, date, ts_format):
    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        return []
    header = [h.lower() for h in rows[0]]

    def col(test):
        return next((i for i, h in enumerate(header) if test(h)), None)

    c = {"ts": col(lambda h: "timestamp" in h),
         "name": col(lambda h: "name" in h and "company" not in h and "exam" not in h),
         "phone": col(lambda h: "phone" in h or "mobile" in h or "contact" in h),
         "email": col(lambda h: "email" in h),
         "company": col(lambda h: "company" in h),
         "designation": col(lambda h: "designation" in h),
         "experience": col(lambda h: "let others know" in h or "experience" in h)}
    log(f"Feedback sheet columns found: { {k: rows[0][v] for k, v in c.items() if v is not None} }")
    for need in ("ts", "company", "designation", "experience"):
        if c[need] is None:
            raise RuntimeError(f"Column for '{need}' not found in the feedback sheet")

    def get(r, key):
        i = c[key]
        return r[i].strip() if i is not None and i < len(r) else ""

    formats = [ts_format] if ts_format else None
    if not formats:
        order = detect_ts_format(get(r, "ts") for r in rows[1:])
        if order:  # with or without the time
            formats = [order + " %H:%M:%S", order]
            log(f"Feedback timestamps are {'month/day' if order.startswith('%m') else 'day/month'}")
    out = [{k: get(r, k) for k in c} for r in rows[1:] if date in timestamp_dates(get(r, "ts"), formats)]
    log(f"{len(out)} feedback response(s) on {date:%d-%b-%Y}")
    return out


def fetch_feedback(cfg, date):
    url = cfg.get("feedback", "sheet_url", fallback="").strip() or cfg.get("feedback", "drive_url")
    driver = open_chrome(cfg, "feedback", url, "google.com")
    if "accounts.google.com" in driver.current_url:
        ask_user("Sign in to Google (advpromktg@gmail.com) in the Ranjit Chrome window.")
    url = find_feedback_sheet(driver, cfg)
    csv_path = download_feedback_csv(driver, url)
    return read_feedback(csv_path, date, cfg.get("feedback", "timestamp_format", fallback="").strip())


# ---------------------------------------------------------------- matching

def name_tokens(s):
    return re.findall(r"[a-z]+", str(s or "").lower())


def name_score(fb_name, row_name):
    """How well a feedback name ("Amit Chowdary", "V.M.Hemanth") fits a sheet name
    ("Doppalapudi Venkata Amit Chowdary", "HEMANTH V M"). 0 = no match, higher = better."""
    f, r = name_tokens(fb_name), name_tokens(row_name)
    if not f or not r or not any(len(t) >= 3 for t in f):
        return 0

    def tok_ok(t):
        if len(t) == 1:  # initial
            return any(x.startswith(t) for x in r)
        return any(t == x or (len(t) >= 3 and len(x) >= 3 and (x.startswith(t) or t.startswith(x)))
                   for x in r)

    if all(tok_ok(t) for t in f):
        return 100 + sum(len(t) for t in f if len(t) >= 3)
    fj, rj = "".join(f), "".join(r)
    if len(fj) >= 5 and (fj in rj or rj in fj):  # "vairamuthu" vs "vaira muthu"
        return 90
    ratio = SequenceMatcher(None, " ".join(sorted(f)), " ".join(sorted(r))).ratio()
    return int(ratio * 80) if ratio >= 0.85 else 0  # small typos


def feedback_columns(ws, hr):
    """Header columns; 'remark' = the Remark right after 'Advantage Pro Feedback form'."""
    cols, after_feedback = {}, False
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(hr, c).value or "").strip().lower()
        if not h:
            continue
        if "feedback" in h and "remark" not in h:
            cols.setdefault("feedback", c)
            after_feedback = True
        elif h == "remark" and after_feedback and "remark" not in cols:
            cols["remark"] = c
        cols.setdefault(h, c)
    cols["company"] = cols.get("company name")
    for need in ("date", "name", "phone no", "email id", "feedback", "remark", "company", "designation"):
        if not cols.get(need):
            raise RuntimeError(f"Column '{need}' not found in sheet {ws.title}")
    return cols


# ---------------------------------------------------------------- excel

def update_excel(path, date, feedback, backup_dir, dry_run=False):
    wb = openpyxl.load_workbook(path)
    ws = month_sheet(wb, date)
    hr = header_row(ws)
    cols = feedback_columns(ws, hr)
    last = last_data_row(ws, hr)
    month_rows = [r for r in range(hr + 1, last + 1) if not is_blank(ws.cell(r, cols["name"]).value)]
    day_rows = [r for r in month_rows if cell_date(ws.cell(r, cols["date"]).value) == date]

    def contact_match(fb, rows):
        p = phone10(fb["phone"])
        email = fb["email"].strip().lower()
        for r in rows:
            if p and p == phone10(ws.cell(r, cols["phone no"]).value):
                return r, "phone"
            if email and email == str(ws.cell(r, cols["email id"]).value or "").strip().lower():
                return r, "email"
        return None, ""

    def name_match(fb):
        scored = sorted(((name_score(fb["name"], ws.cell(r, cols["name"]).value), r) for r in day_rows),
                        reverse=True)
        if not scored or scored[0][0] == 0:
            return None, ""
        if len(scored) > 1 and scored[1][0] == scored[0][0]:
            names = [clean_cell_text(ws.cell(r, cols["name"]).value) for s, r in scored if s == scored[0][0]]
            log(f"  ? '{fb['name']}' fits more than one row ({', '.join(names)}) - fill it by hand")
            return None, ""
        return scored[0][1], "name"

    updated, unmatched = 0, []
    for fb in feedback:
        row, how = contact_match(fb, day_rows)
        if not row:
            row, how = name_match(fb)
        if not row:
            row, how = contact_match(fb, month_rows)
        if not row:
            unmatched.append(fb)
            continue

        changed = []

        def fill(col, value):
            if value and is_blank(ws.cell(row, cols[col]).value):
                ws.cell(row, cols[col]).value = value
                changed.append(col)

        fill("company", clean_cell_text(fb["company"]))
        fill("designation", clean_cell_text(fb["designation"]))
        fill("feedback", "y")
        fill("remark", clean_cell_text(fb["experience"]))
        who = clean_cell_text(ws.cell(row, cols["name"]).value)
        if changed:
            updated += 1
            log(f"  row {row} {who} <- feedback '{fb['name']}' (matched by {how}): "
                f"company '{clean_cell_text(fb['company'])}' | designation '{clean_cell_text(fb['designation'])}' "
                f"| remark '{clean_cell_text(fb['experience'])[:60]}'")
        else:
            log(f"  row {row} {who}: feedback already filled")

    if unmatched:
        log(f"{len(unmatched)} response(s) not in the Red Hat sheet (Pearson VUE or other exams):")
        for fb in unmatched:
            log(f"    - {fb['name']} | {fb['phone']} | {fb['email']}")
    log(f"Sheet {ws.title}: {updated} row(s) updated with feedback")

    if dry_run:
        log("Dry run - Excel file NOT saved")
        return
    if not updated:
        log("Nothing changed - Excel file not saved")
        return
    bdir = HERE / backup_dir
    bdir.mkdir(exist_ok=True)
    backup = bdir / f"{Path(path).stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.copy2(path, backup)
    log(f"Backup saved: {backup}")
    wb.save(path)
    log("Excel file saved")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", help="YYYY-MM-DD (default: asks, Enter = today)")
    ap.add_argument("--dry-run", action="store_true", help="do everything except saving the Excel file")
    args = ap.parse_args()

    cfg = load_config()
    date = dt.date.fromisoformat(args.date) if args.date else ask_date()
    log(f"Date: {date:%d-%b-%Y}   Excel: {cfg.get('excel', 'path')}")
    run(cfg, date, dry_run=args.dry_run)
    log("Done")


def run(cfg, date, dry_run=False):
    """Feedback responses of `date` -> company, designation, feedback y, remark."""
    excel = cfg.get("excel", "path")
    check_excel_closed(excel)

    feedback = fetch_feedback(cfg, date)
    if not feedback:
        log("No feedback that day - nothing to write")
        return
    check_excel_closed(excel)
    update_excel(excel, date, feedback, cfg.get("excel", "backup_dir"), dry_run=dry_run)


if __name__ == "__main__":
    main()
