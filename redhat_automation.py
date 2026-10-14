"""
Red Hat Exam Scheduler + Adv Pro  ->  "redhat exam 2026.xlsx"

Steps
  1. Open Chrome (Advantage pro profile), Red Hat Exam Scheduler > menu >
     Exam Deliveries. Status = Completed, Time Zone = Asia/Calcutta,
     Effective Date From/To = the day, Location = 2093, then 2337, then 2394.
  2. For every candidate: name, email, contact number, exam code (first 5
     characters, e.g. EX200) and "retake" if the eligibility code says RETAKE,
     otherwise "y".
  3. Open Adv Pro (vetri), Master > Registration > Display, filter by
     PresentPhoneNo and search every phone number: found = "y" (our student),
     not found = "n".
  4. Back up the Excel file, add the rows to the month sheet, save.

Usage
  python redhat_automation.py                    # asks for the date (Enter = today)
  python redhat_automation.py --date 2026-10-08  # a specific day
  python redhat_automation.py --skip-advpro      # don't check Adv Pro ("Our Student" left blank)
  python redhat_automation.py --dry-run          # don't save the Excel file
"""
import argparse
import configparser
import datetime as dt
import json
import re
import shutil
import socket
import subprocess
import sys
import time
import urllib.request
import warnings
from copy import copy
from pathlib import Path

import openpyxl
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys

HERE = Path(__file__).resolve().parent
CONFIG_PATH = HERE / "config.ini"
MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
UPPER, LOWER = "ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz"


# ---------------------------------------------------------------- helpers

def log(msg):
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)


def ask_user(msg):
    input(f"\n>>> {msg}\n    Press Enter when done... ")


def parse_date(text):
    """'08-10-2026', '8/10/2026', '8.10.26', '2026-10-08' or '8-10' (this year). Day comes first."""
    text = text.strip()
    if re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", text):
        return dt.date.fromisoformat("-".join(p.zfill(2) for p in text.split("-")))
    parts = re.split(r"[-/.\s]+", text)
    if not 2 <= len(parts) <= 3 or not all(p.isdigit() for p in parts):
        raise ValueError(text)
    day, month = int(parts[0]), int(parts[1])
    year = int(parts[2]) if len(parts) == 3 else dt.date.today().year
    if year < 100:
        year += 2000
    return dt.date(year, month, day)


def ask_date():
    today = dt.date.today()
    while True:
        text = input(f"\n>>> Which date? Type it as DD-MM-YYYY (e.g. {today:%d-%m-%Y}),\n"
                     f"    or just press Enter for today ({today:%d-%b-%Y}): ").strip(" \t﻿")
        if not text:
            return today
        try:
            date = parse_date(text)
        except ValueError:
            print("    Not a valid date - try again, e.g. 07-10-2026")
            continue
        print(f"    Using {date:%A, %d-%b-%Y}")
        return date


def load_config():
    cfg = configparser.RawConfigParser(inline_comment_prefixes=(";",))
    cfg.read(CONFIG_PATH, encoding="utf-8")
    return cfg


def norm_name(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def clean_cell_text(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def phone10(value):
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-10:] if len(digits) >= 10 else ""


def clean_phone(raw):
    """Indian mobile -> int like the existing sheet; anything else stays as text."""
    raw = clean_cell_text(raw)
    if not raw:
        return None
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) == 10 and digits[0] in "6789":
        return int(digits)
    return raw


def is_blank(value):
    return value is None or str(value).strip() == ""


def exam_code(text):
    """'EX200V10K: Red Hat Certified System Administrator' -> 'EX200'."""
    text = clean_cell_text(text)
    m = re.match(r"[A-Za-z]{2}\d{3}", text)
    return (m.group(0) if m else text[:5]).upper()


# ---------------------------------------------------------------- chrome

def port_open(port):
    with socket.socket() as s:
        s.settimeout(0.5)
        return s.connect_ex(("127.0.0.1", port)) == 0


def shortcut_profile(name):
    """Read --profile-directory from a Desktop shortcut like 'Advantage pro - Chrome.lnk'."""
    for desktop in (Path.home() / "Desktop", Path(r"C:\Users\Public\Desktop")):
        lnk = desktop / name
        if lnk.exists():
            args = subprocess.run(
                ["powershell", "-NoProfile", "-Command",
                 f"(New-Object -ComObject WScript.Shell).CreateShortcut('{lnk}').Arguments"],
                capture_output=True, text=True).stdout
            m = re.search(r'--profile-directory=(?:"([^"]+)"|(\S+))', args)
            if m:
                return m.group(1) or m.group(2)
            raise RuntimeError(f"No --profile-directory in shortcut {lnk}")
    raise RuntimeError(f"Shortcut not found on the Desktop: {name}")


def chrome_running():
    out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq chrome.exe"], capture_output=True, text=True).stdout
    return "chrome.exe" in out.lower()


# keep Chrome from freezing/throttling tabs in the background
CHROME_FLAGS = ["--disable-background-timer-throttling", "--disable-renderer-backgrounding",
                "--disable-backgrounding-occluded-windows",
                "--disable-features=IntensiveWakeUpThrottling,CalculateNativeWinOcclusion"]


def chrome_pages(port):
    """Tabs of the debug-port Chrome, straight from Chrome (never hangs on a frozen tab)."""
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/json/list", timeout=10) as r:
            return [t for t in json.load(r) if t.get("type") == "page"]
    except Exception:
        return []


def activate_tab(port, tab_id):
    try:
        urllib.request.urlopen(f"http://127.0.0.1:{port}/json/activate/{tab_id}", timeout=10).close()
    except Exception:
        pass


def wake_tabs(port):
    """Chrome freezes background tabs and ChromeDriver hangs on a frozen tab;
    bringing each tab to the front for a moment wakes it up."""
    for t in chrome_pages(port):
        if t["url"].startswith("http"):
            activate_tab(port, t["id"])
            time.sleep(0.3)


def open_chrome(cfg, section, start_url, match):
    """Open start_url in the Chrome profile of the section's Desktop shortcut and
    return the Selenium driver switched to that tab.

    Chrome is started with a debug port so the script can control it. That only
    works if Chrome was fully closed before the first launch.
    """
    port = cfg.getint("chrome", "debug_port")
    profile = shortcut_profile(cfg.get(section, "shortcut"))
    if not port_open(port) and chrome_running():
        # Chrome often keeps running in the background after its windows are closed
        answer = input("\n>>> Chrome is already running (maybe in the background).\n"
                       "    The script must restart it. Close Chrome now? Open tabs come back\n"
                       "    with Ctrl+Shift+T. [Y/n]: ").strip().lower()
        if answer in ("", "y", "yes"):
            subprocess.run(["taskkill", "/IM", "chrome.exe", "/F"], capture_output=True)
            for _ in range(20):
                if not chrome_running():
                    break
                time.sleep(0.5)
        if chrome_running():
            sys.exit("Chrome is still running - close it (check the system tray) and run again.")
        log("Chrome closed")
    before = {t["id"] for t in chrome_pages(port)}
    log(f"Opening Chrome profile '{profile}' ({cfg.get(section, 'shortcut')})")
    subprocess.Popen([cfg.get("chrome", "path"), f"--profile-directory={profile}",
                      f"--remote-debugging-port={port}", *CHROME_FLAGS, start_url])
    for _ in range(60):
        if port_open(port):
            break
        time.sleep(0.5)
    else:
        raise RuntimeError(f"Chrome did not open the control port {port}")

    # find the tab that just opened start_url (newest first) without touching the others
    target = None
    end = time.time() + 30
    while time.time() < end and not target:
        pages = [t for t in chrome_pages(port) if match in t["url"]]
        new = [t for t in pages if t["id"] not in before]
        target = (new or pages or [None])[0]
        if not target:
            time.sleep(1)
    if not target:
        ask_user(f"Could not find the new Chrome tab ({match}). Click on it in Chrome.")
        target = next((t for t in chrome_pages(port) if match in t["url"]), None)
        if not target:
            raise RuntimeError(f"No Chrome tab with {match}")

    log("Waking up Chrome tabs")
    wake_tabs(port)
    activate_tab(port, target["id"])
    opts = Options()
    opts.add_experimental_option("debuggerAddress", f"127.0.0.1:{port}")
    driver = webdriver.Chrome(options=opts)
    log("Connected to Chrome")
    driver.switch_to.window(target["id"])
    return driver


def visible(el):
    try:
        return el.is_displayed()
    except Exception:
        return False


def click(driver, el):
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)


def has_text(driver, text):
    xp = f"//*[normalize-space(translate(., '{UPPER}', '{LOWER}'))='{text.lower()}']"
    return any(visible(e) for e in driver.find_elements(By.XPATH, xp))


def click_text(driver, *texts, timeout=10):
    """Click the visible element whose whole text equals one of `texts` (case-insensitive)."""
    end = time.time() + timeout
    while True:
        for text in texts:
            xp = f"//*[normalize-space(translate(., '{UPPER}', '{LOWER}'))='{text.lower()}']"
            els = [e for e in driver.find_elements(By.XPATH, xp) if visible(e)]
            if els:
                click(driver, els[-1])  # last = innermost element
                return True
        if time.time() > end:
            return False
        time.sleep(0.5)


def wait_for(driver, css, timeout=20):
    end = time.time() + timeout
    while time.time() < end:
        els = [e for e in driver.find_elements(By.CSS_SELECTOR, css) if visible(e)]
        if els:
            return els[0]
        time.sleep(0.5)
    return None


def type_into(el, value):
    el.click()
    el.send_keys(Keys.CONTROL, "a")
    el.send_keys(Keys.DELETE)
    el.send_keys(value)


# ---------------------------------------------------------------- red hat

STATUS = "md-select[formcontrolname=searchStatusControl]"
TIMEZONE = "input[formcontrolname=searchTimeZoneControl]"
LOCATION = "input[formcontrolname=searchLocationControl]"
DATE_FROM = "input[formcontrolname=searchFromDateControl]"
DATE_TO = "input[formcontrolname=searchToDateControl]"


def open_exam_deliveries(driver):
    end = time.time() + 30
    while time.time() < end and not driver.find_elements(By.CSS_SELECTOR, ".menuIconCls"):
        if "sso.redhat.com" in driver.current_url or "login" in driver.current_url.lower():
            ask_user("Please log in to Red Hat in the Chrome window.")
            end = time.time() + 30
        time.sleep(1)
    menu = wait_for(driver, ".menuIconCls", timeout=5)
    if menu:
        click(driver, menu)
        time.sleep(1)
    if click_text(driver, "Exam Deliveries", timeout=10):
        log("Clicked menu > Exam Deliveries")
    if not wait_for(driver, STATUS, timeout=20):
        ask_user("Could not open 'Exam Deliveries'. Click the 3 lines (top left) > Exam Deliveries.")
    time.sleep(1)


def pick_option(driver, text, timeout=10):
    """Click an open dropdown/autocomplete option containing `text`."""
    end = time.time() + timeout
    while time.time() < end:
        for o in driver.find_elements(By.CSS_SELECTOR, "md-option, mat-option"):
            if visible(o) and text.lower() in clean_cell_text(o.text).lower():
                click(driver, o)
                return True
        time.sleep(0.5)
    return False


def set_status(driver, status):
    sel = wait_for(driver, STATUS)
    if clean_cell_text(sel.text).lower().endswith(status.lower()):
        return
    click(driver, sel)
    time.sleep(0.5)
    if not pick_option(driver, status):
        driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
        ask_user(f"Could not choose Status '{status}'. Please choose it in Chrome.")
    log(f"Status = {status}")


def set_autocomplete(driver, css, typed, option_text, label):
    el = wait_for(driver, css)
    type_into(el, typed)
    time.sleep(1)
    if not pick_option(driver, option_text):
        ask_user(f"Could not choose {label} '{option_text}'. Please choose it in Chrome.")
    time.sleep(0.5)
    log(f"{label} = {clean_cell_text(el.get_attribute('value'))}")


def set_date(driver, css, date, label):
    el = wait_for(driver, css)
    type_into(el, date.isoformat())
    el.send_keys(Keys.TAB)
    time.sleep(0.5)
    value = el.get_attribute("value") or ""
    if value != date.isoformat():
        ask_user(f"{label} date shows '{value}'. Please set it to {date:%d-%b-%Y} in Chrome.")
    log(f"{label} = {date.isoformat()}")


READ_ROWS_JS = r"""
const vis = e => !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
const cell = (r, k) => { const c = r.querySelector('.cdk-column-' + k); return c ? c.innerText.replace(/\s+/g, ' ').trim() : ''; };
const keys = ['candidateName', 'candidateEmail', 'candidateContactNumber', 'examCodeName',
              'eligibilityCode', 'locationName', 'examStartDateTime', 'statusDisplay'];
const rows = [...document.querySelectorAll('md-row, mat-row')].filter(vis)
    .map(r => Object.fromEntries(keys.map(k => [k, cell(r, k)])));
const label = document.querySelector('.mat-paginator-range-label');
const next = document.querySelector('.mat-paginator-navigation-next');
const busy = [...document.querySelectorAll('md-spinner, md-progress-spinner, md-progress-bar, mat-spinner, .loading, .spinner')].some(vis);
return {rows, label: label ? label.innerText.trim() : '', hasNext: !!(next && !next.disabled), busy};
"""


# Count the page's server requests, so after SEARCH / Next we can wait for the
# answer instead of reading the old table.
TRACK_REQUESTS_JS = r"""
if (!window.__rq) {
  window.__rq = {pending: 0, done: 0};
  const send = XMLHttpRequest.prototype.send;
  XMLHttpRequest.prototype.send = function () {
    window.__rq.pending++;
    this.addEventListener('loadend', () => { window.__rq.pending--; window.__rq.done++; });
    return send.apply(this, arguments);
  };
  if (window.fetch) {
    const f = window.fetch;
    window.fetch = function () {
      window.__rq.pending++;
      return f.apply(this, arguments).finally(() => { window.__rq.pending--; window.__rq.done++; });
    };
  }
}
return window.__rq.done;
"""


def requests_done(driver):
    return driver.execute_script(TRACK_REQUESTS_JS)


def read_page(driver, done_before=None, must_ask=True, settle=1.0, timeout=30):
    """Wait for the server answer (if done_before is given) and until the table
    stops changing, then return its rows. must_ask=False: the page may sort
    things out by itself (paging), so don't wait long if it asks nothing."""
    start = time.time()
    end = start + timeout
    if done_before is not None:
        while time.time() < end:
            rq = driver.execute_script("return window.__rq || {pending: 0, done: 0};")
            if rq["pending"] == 0 and (rq["done"] > done_before or (not must_ask and time.time() - start > 3)):
                break
            time.sleep(0.3)
        else:
            log("  (the page did not answer in time - reading the table anyway)")
    last = None
    while True:
        time.sleep(settle)
        page = driver.execute_script(READ_ROWS_JS)
        key = (page["label"], json.dumps(page["rows"]))
        if not page["busy"] and key == last:
            return page
        if time.time() > end:
            return page
        last = key


def biggest_page_size(driver):
    sel = wait_for(driver, ".mat-paginator-page-size-select", timeout=5)
    if not sel:
        return
    click(driver, sel)
    time.sleep(0.5)
    opts = [o for o in driver.find_elements(By.CSS_SELECTOR, "md-option, mat-option") if visible(o)]
    if opts:
        click(driver, opts[-1])
        time.sleep(1)
    else:
        driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)


def button_text(b):
    # buttons start with an icon-font character (e.g.  before SEARCH)
    return clean_cell_text(re.sub(r"[-]", "", b.text)).upper()


def click_button(driver, text, timeout=5):
    end = time.time() + timeout
    while time.time() < end:
        for b in driver.find_elements(By.TAG_NAME, "button"):
            if visible(b) and button_text(b) == text.upper():
                click(driver, b)
                return True
        time.sleep(0.5)
    log(f"  buttons on the page: {[(clean_cell_text(b.text), visible(b)) for b in driver.find_elements(By.TAG_NAME, 'button')]}")
    return False


def search_location(driver, cfg, date, code):
    set_autocomplete(driver, LOCATION, code, f"- {code}", "Location")
    before = requests_done(driver)
    if not click_button(driver, "SEARCH"):
        ask_user("Could not click SEARCH. Please click it in Chrome.")
        before = None
    log(f"Searching location {code}")
    page = read_page(driver, before)
    if page["hasNext"]:
        before = requests_done(driver)
        biggest_page_size(driver)
        page = read_page(driver, before, must_ask=False)

    rows, seen_pages = list(page["rows"]), {page["label"]}
    while page["hasNext"]:
        before = requests_done(driver)
        nxt = driver.find_element(By.CSS_SELECTOR, ".mat-paginator-navigation-next")
        click(driver, nxt)
        page = read_page(driver, before, must_ask=False)
        if page["label"] in seen_pages:
            break
        seen_pages.add(page["label"])
        rows += page["rows"]
    log(f"  {page['label'] or len(rows)} result(s)")

    people = []
    status = cfg.get("redhat", "status").lower()
    for r in rows:
        name = clean_cell_text(r["candidateName"])
        # safety: only that location, that day and that status
        if code not in r["locationName"]:
            log(f"  skipped (other location): {name} | {r['locationName']}")
            continue
        if not r["examStartDateTime"].startswith(date.isoformat()):
            log(f"  skipped (other date): {name} | {r['examStartDateTime']}")
            continue
        if r["statusDisplay"].lower() != status:
            log(f"  skipped (status {r['statusDisplay']}): {name}")
            continue
        person = {"location": code, "name": name, "email": clean_cell_text(r["candidateEmail"]),
                  "phone": clean_cell_text(r["candidateContactNumber"]),
                  "code": exam_code(r["examCodeName"]),
                  "status": "retake" if "retake" in r["eligibilityCode"].lower() else "y",
                  "eligibility": r["eligibilityCode"], "our_student": ""}
        people.append(person)
        log(f"  {name} | {person['phone']} | {person['email']} | {person['code']} | "
            f"{person['eligibility']} -> {person['status']}")
    return people


def scrape_redhat(cfg, date):
    driver = open_chrome(cfg, "redhat", cfg.get("redhat", "url"), "rhtapps.redhat.com")
    open_exam_deliveries(driver)
    set_status(driver, cfg.get("redhat", "status"))
    tz = cfg.get("redhat", "time_zone")
    set_autocomplete(driver, TIMEZONE, tz, tz, "Time Zone")
    set_date(driver, DATE_FROM, date, "From")
    set_date(driver, DATE_TO, date, "To")
    people = []
    for code in [c.strip() for c in cfg.get("redhat", "locations").split(",") if c.strip()]:
        people += search_location(driver, cfg, date, code)
    log(f"{len(people)} completed exam(s) on {date:%d-%b-%Y}")
    return people


# ---------------------------------------------------------------- adv pro

class AdvPro:
    """Drives the Adv Pro desktop program: login, Master > Registration > Display,
    Filter By PresentPhoneNo, and tells if a phone number is registered."""

    GRID = "DataGridWndClass"

    def __init__(self, cfg):
        from pywinauto import Application
        # pywinauto turns its own warnings back on when imported
        warnings.filterwarnings("ignore", message=".*32-bit application.*")
        self.cfg = cfg
        exe = cfg.get("advpro", "exe")
        try:
            self.app = Application(backend="win32").connect(path=Path(exe).name, timeout=2)
            log("Adv Pro is already open")
        except Exception:
            log("Opening Adv Pro")
            self.app = Application(backend="win32").start(f'"{exe}"', work_dir=str(Path(exe).parent))
        self.main = self.app.window(class_name="ThunderRT6MDIForm")
        self.login()
        self.form = self.open_display()
        self.grid = self.form.child_window(class_name=self.GRID).wrapper_object()
        self.search = self.form.child_window(class_name="ThunderRT6TextBox", control_id=51).wrapper_object()
        combo = self.form.child_window(class_name="ThunderRT6ComboBox", control_id=52).wrapper_object()
        combo.select("PresentPhoneNo")
        time.sleep(0.5)
        if combo.selected_text() != "PresentPhoneNo":
            ask_user("Please choose 'PresentPhoneNo' in 'Filter By' in Adv Pro.")
        log("Adv Pro: Filter By = PresentPhoneNo")

    def login(self):
        login = self.app.window(title="Form1", class_name="ThunderRT6FormDC")
        end = time.time() + 30
        while time.time() < end:
            if self.main.exists() and self.main.is_visible():
                return
            if login.exists() and login.is_visible():
                break
            time.sleep(1)
        else:
            ask_user("Adv Pro did not show the login screen. Log in by hand.")
            return
        user = login.child_window(class_name="ThunderRT6TextBox", control_id=1).wrapper_object()
        pwd = login.child_window(class_name="ThunderRT6TextBox", control_id=2).wrapper_object()
        self.set_text(user, self.cfg.get("advpro", "username"))
        self.set_text(pwd, self.cfg.get("advpro", "password"))
        pwd.set_focus()
        pwd.type_keys("{ENTER}")
        if not self.main.wait("visible", timeout=30):
            ask_user("Adv Pro login did not finish. Log in by hand.")
        log("Adv Pro: logged in")

    def display_form(self):
        for child in self.main.children(class_name="ThunderRT6FormDC"):
            if child.is_visible() and child.children(class_name=self.GRID):
                return self.main.child_window(handle=child.handle)
        return None

    def open_display(self):
        form = self.display_form()
        if not form:
            self.main.set_focus()
            self.main.menu_select("Master->Registration->Display")
            log("Adv Pro: Master > Registration > Display")
            end = time.time() + 20
            while not form and time.time() < end:
                time.sleep(1)
                form = self.display_form()
        if not form:
            ask_user("Open Master > Registration > Display in Adv Pro.")
            form = self.display_form()
        if not form:
            raise RuntimeError("Adv Pro Registration Display screen not found")
        return form

    @staticmethod
    def set_text(ctrl, text):
        import win32con
        import win32gui
        win32gui.SendMessage(ctrl.handle, win32con.WM_SETTEXT, 0, text)

    def grid_has_rows(self):
        """The list shows columns (and a scroll bar) only when something was found.
        Read it two ways: UI Automation, and the screen pixels of the list."""
        from pywinauto.controls.uiawrapper import UIAWrapper
        from pywinauto.uia_element_info import UIAElementInfo
        try:
            uia = len(UIAWrapper(UIAElementInfo(self.grid.handle)).children()) > 0
        except Exception:
            uia = None
        try:
            from PIL import ImageGrab
            r = self.grid.rectangle()
            img = ImageGrab.grab(bbox=(r.left, r.top, r.right, r.bottom), all_screens=True).convert("RGB")
            w, h = img.size
            px = img.load()
            # an empty list is grey right of the first two blank columns
            white = sum(1 for x in range(260, w - 4, 2) for y in range(2, h - 18, 2) if min(px[x, y]) > 245)
            pixels = white > 150
        except Exception:
            pixels = None
        return uia, pixels

    def is_student(self, phone):
        p = phone10(phone)
        if not p:
            return ""
        self.set_text(self.search, p)
        wait = self.cfg.getfloat("advpro", "search_wait", fallback=2)
        time.sleep(wait)
        last = self.grid_has_rows()
        for _ in range(6):  # wait until the list stops changing
            time.sleep(0.7)
            now = self.grid_has_rows()
            if now == last:
                break
            last = now
        uia, pixels = last
        if uia is not None and pixels is not None and uia != pixels:
            log(f"  ? Adv Pro {p}: unsure (list check {uia}, screen check {pixels}) - check by hand")
        found = uia if uia is not None else pixels
        return "y" if found else "n"


def check_advpro(cfg, people):
    todo = [p for p in people if phone10(p["phone"])]
    if not todo:
        return
    adv = AdvPro(cfg)
    adv.main.set_focus()  # keep it in front for the screen check
    for p in people:
        if not phone10(p["phone"]):
            log(f"  Adv Pro: {p['name']} has no phone number - 'Our Student' left blank")
            continue
        p["our_student"] = adv.is_student(p["phone"])
        log(f"  Adv Pro {phone10(p['phone'])} {p['name']}: {p['our_student']}")


# ---------------------------------------------------------------- excel

def check_excel_closed(path):
    try:
        with open(path, "r+b"):
            pass
    except PermissionError:
        sys.exit(f"The Excel file is open on some computer. Close it and run again:\n  {path}")


def month_sheet(wb, date):
    abbr = MONTHS[date.month - 1]
    for ws in wb.worksheets:
        if ws.title.strip().upper().startswith(abbr):
            return ws
    raise RuntimeError(f"No sheet for {abbr} in the Excel file - add it first (copy last month's).")


def header_row(ws):
    """The candidate table header: 'Date' in A and 'Name' in B (the count table above has 2093 in B)."""
    for r in range(1, ws.max_row + 1):
        if (str(ws.cell(r, 1).value or "").strip().lower() == "date"
                and str(ws.cell(r, 2).value or "").strip().lower() == "name"):
            return r
    raise RuntimeError(f"Header row (Date, Name) not found in sheet {ws.title}")


def sheet_columns(ws, hr):
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(hr, c).value or "").strip().lower()
        if h:
            cols.setdefault(h, c)
    for need in ("date", "name", "phone no", "email id", "exam code", "exam status", "our student y/n"):
        if need not in cols:
            raise RuntimeError(f"Column '{need}' not found in sheet {ws.title}")
    return cols


def cell_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def last_data_row(ws, hr):
    for r in range(ws.max_row, hr, -1):
        if not is_blank(ws.cell(r, 1).value) or not is_blank(ws.cell(r, 2).value):
            return r
    return hr


def update_excel(path, date, people, backup_dir, dry_run=False):
    wb = openpyxl.load_workbook(path)
    ws = month_sheet(wb, date)
    hr = header_row(ws)
    cols = sheet_columns(ws, hr)

    def day_rows():
        return [r for r in range(hr + 1, last_data_row(ws, hr) + 1)
                if cell_date(ws.cell(r, cols["date"]).value) == date]

    def same_person(r, p):
        email = p["email"].lower()
        return ((email and email == str(ws.cell(r, cols["email id"]).value or "").strip().lower())
                or (phone10(p["phone"]) and phone10(p["phone"]) == phone10(ws.cell(r, cols["phone no"]).value))
                or norm_name(p["name"]) == norm_name(ws.cell(r, cols["name"]).value))

    added = 0
    for p in people:
        rows = day_rows()
        if any(same_person(r, p) for r in rows if not is_blank(ws.cell(r, cols["name"]).value)):
            log(f"  already in sheet: {p['name']}")
            continue
        # use the empty row already made for this date, else add one after that date's rows
        blank = [r for r in rows if is_blank(ws.cell(r, cols["name"]).value)]
        if blank:
            row = blank[0]
        else:
            after = rows[-1] if rows else max(
                [r for r in range(hr + 1, last_data_row(ws, hr) + 1)
                 if (cell_date(ws.cell(r, cols["date"]).value) or dt.date.min) < date] or [hr])
            row = after + 1
            if row <= ws.max_row:
                ws.insert_rows(row)
            style_row = row - 1 if row - 1 > hr else row + 1
            for c in range(1, ws.max_column + 1):
                ws.cell(row, c)._style = copy(ws.cell(style_row, c)._style)
            ws.row_dimensions[row].height = ws.row_dimensions[style_row].height
        ws.cell(row, cols["date"]).value = dt.datetime(date.year, date.month, date.day)
        ws.cell(row, cols["name"]).value = p["name"]
        ws.cell(row, cols["phone no"]).value = clean_phone(p["phone"])
        ws.cell(row, cols["email id"]).value = p["email"] or None
        ws.cell(row, cols["exam code"]).value = p["code"]
        ws.cell(row, cols["exam status"]).value = p["status"]
        ws.cell(row, cols["our student y/n"]).value = p["our_student"] or None
        added += 1
        log(f"  row {row}: {p['name']} | {p['phone']} | {p['email']} | {p['code']} | "
            f"{p['status']} | our student {p['our_student'] or '-'}  ({p['location']})")

    log(f"Sheet {ws.title}: {added} row(s) added")
    if dry_run:
        log("Dry run - Excel file NOT saved")
        return
    if not added:
        log("Nothing changed - Excel file not saved")
        return
    bdir = HERE / backup_dir
    bdir.mkdir(exist_ok=True)
    backup = bdir / f"{Path(path).stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.copy2(path, backup)
    log(f"Backup saved: {backup}")
    wb.save(path)
    log("Excel file saved")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", help="YYYY-MM-DD (default: asks, Enter = today)")
    ap.add_argument("--skip-advpro", action="store_true", help="don't check the phone numbers in Adv Pro")
    ap.add_argument("--dry-run", action="store_true", help="do everything except saving the Excel file")
    args = ap.parse_args()

    cfg = load_config()
    date = dt.date.fromisoformat(args.date) if args.date else ask_date()
    log(f"Date: {date:%d-%b-%Y}   Excel: {cfg.get('excel', 'path')}")
    run(cfg, date, skip_advpro=args.skip_advpro, dry_run=args.dry_run)
    log("Done")


def run(cfg, date, skip_advpro=False, dry_run=False):
    """Red Hat completed exams of `date` + Adv Pro check -> new rows in the Excel file."""
    excel = cfg.get("excel", "path")
    check_excel_closed(excel)

    people = scrape_redhat(cfg, date)
    if people and not skip_advpro:
        check_advpro(cfg, people)

    logs = HERE / "logs"
    logs.mkdir(exist_ok=True)
    (logs / f"{date}_{dt.datetime.now():%H%M%S}.json").write_text(
        json.dumps(people, indent=2, ensure_ascii=False), encoding="utf-8")

    if not people:
        log("No completed exams found - nothing to write")
        return
    check_excel_closed(excel)
    update_excel(excel, date, people, cfg.get("excel", "backup_dir"), dry_run=dry_run)


if __name__ == "__main__":
    main()
